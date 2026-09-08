"""XLSX conversion for the client deliverable and its return trip (POC-013).

This module converts the CSV export *without modifying it*: same 16 columns of
PROFILE_CSV_FIELDS, same order, same values. It adds a format, not a view.

Why it exists (call of 04/09/2026, action item no 1): the client must be able
to sort and filter the deliverable in Excel, and to send it back annotated
without being asked to convert it. Delivering xlsx while demanding CSV on the
way back would move the conversion chore to a French Excel, whose "save as
CSV" produces a semicolon-separated, BOM-less file that our reader - which
expects commas - would mis-parse.

The real risk of this module is not the code, it is what Excel does to the
data in transit. csv.DictReader yields str; openpyxl yields int, float,
datetime and None. run_poc006 runs enregistrer_profils *before*
importer_commentaires_csv, so a value reformatted by Excel would be written
straight into the store by that first pass. Hence the two guardrails below:

* on write, every cell is a text cell (explicit "@" number format), so Excel
  neither reinterprets a date nor turns a score into a float;
* on read, every value is brought back to str and an empty cell to "", so the
  comparisons in importer_commentaires_csv and _importer_statut_coordonnees
  behave exactly as they do on a CSV.

A header that no longer matches the export is reported and blocking, never
guessed - same requirement as the statuts_refuses of POC-009: business data is
expensive, we do not fix it silently.

No Streamlit import here (see CLAUDE.md / ARCHITECTURE.md).
"""

import csv
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from source.backend.adapters.storage.csv_export import PROFILE_CSV_FIELDS

# Client deliverable, sibling of EXPORT_CSV. Gitignored like every data file.
EXPORT_XLSX = Path("./profils_magasin.xlsx")

# Excel's "Text" number format. Applied to every data cell so that a value the
# client retypes stays a string instead of becoming a date or a number.
_FORMAT_TEXTE = "@"


class EnteteXlsxInvalide(ValueError):
    """The workbook's header does not match PROFILE_CSV_FIELDS.

    Raised rather than worked around: a column inserted, removed or reordered
    by the client would silently mis-pair values with columns.
    """


def _cellule_en_texte(valeur: object) -> str:
    """Bring any openpyxl cell value back to the str csv.DictReader would give.

    An empty cell is "" and never None: _importer_statut_coordonnees calls
    .strip().lower() on it. An integral float loses its ".0" so that a score
    read back as 4.0 still compares equal to the stored "4". A datetime at
    midnight is a plain ISO date, which is how date_collecte is stored.
    """
    if valeur is None:
        return ""
    if isinstance(valeur, str):
        return valeur
    # bool before int: bool is a subclass of int, and Excel's TRUE/FALSE must
    # come back lowercase to be recognised by _VALEURS_OPPOSITION.
    if isinstance(valeur, bool):
        return "true" if valeur else "false"
    if isinstance(valeur, int):
        return str(valeur)
    if isinstance(valeur, float):
        return str(int(valeur)) if valeur == int(valeur) else str(valeur)
    if isinstance(valeur, datetime):
        jour = valeur.date().isoformat()
        return jour if valeur.time() == time(0, 0) else f"{jour} {valeur.time().isoformat()}"
    if isinstance(valeur, (date, time)):
        return valeur.isoformat()
    return str(valeur)


def _valider_entete(entete: list[str], source: Path) -> None:
    """Reject a header that is not exactly PROFILE_CSV_FIELDS, in order."""
    if entete == PROFILE_CSV_FIELDS:
        return

    manquantes = [c for c in PROFILE_CSV_FIELDS if c not in entete]
    inconnues = [c for c in entete if c not in PROFILE_CSV_FIELDS]
    if manquantes or inconnues:
        detail = []
        if manquantes:
            detail.append(f"colonnes manquantes : {', '.join(manquantes)}")
        if inconnues:
            detail.append(f"colonnes inconnues : {', '.join(inconnues)}")
        raison = " ; ".join(detail)
    else:
        raison = "colonnes reordonnees"

    raise EnteteXlsxInvalide(
        f"En-tete non conforme dans {source} ({raison}).\n"
        f"  attendu : {PROFILE_CSV_FIELDS}\n"
        f"  lu      : {entete}\n"
        "Le fichier n'est pas relu : apparier des valeurs sur un en-tete "
        "modifie ecrirait des donnees dans les mauvaises colonnes."
    )


def lire_xlsx(xlsx_path: Path) -> list[dict[str, str]]:
    """Read an annotated workbook as csv.DictReader would read the CSV.

    Every value is a str and every empty cell is "". The header must match
    PROFILE_CSV_FIELDS exactly, otherwise EnteteXlsxInvalide is raised and
    nothing is returned.
    """
    classeur = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        feuille = classeur.worksheets[0]
        lignes = feuille.iter_rows(values_only=True)

        try:
            brute = next(lignes)
        except StopIteration:
            raise EnteteXlsxInvalide(
                f"Classeur vide dans {xlsx_path} : aucune ligne d'en-tete."
            ) from None

        entete = [_cellule_en_texte(v).strip() for v in brute]
        # Excel commonly reports trailing empty columns; they carry no name.
        while entete and not entete[-1]:
            entete.pop()
        _valider_entete(entete, xlsx_path)

        profils: list[dict[str, str]] = []
        for brute in lignes:
            valeurs = [_cellule_en_texte(v) for v in brute][: len(entete)]
            # A row Excel padded with blanks is not a profile.
            if not any(valeurs):
                continue
            valeurs += [""] * (len(entete) - len(valeurs))
            profils.append(dict(zip(entete, valeurs)))
        return profils
    finally:
        classeur.close()


def convertir_csv_en_xlsx(csv_path: Path, xlsx_path: Path = EXPORT_XLSX) -> int:
    """Convert the CSV export to a workbook, value for value. Returns the row count.

    Nothing is reformatted: the cells carry the CSV strings as text. The only
    additions are a frozen header row and an auto-filter (user arbitration of
    08/09/2026), which touch no value.
    """
    with open(csv_path, newline="", encoding="utf-8") as f:
        lecteur = csv.DictReader(f)
        entete = list(lecteur.fieldnames or [])
        _valider_entete(entete, csv_path)
        lignes = list(lecteur)

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Profils"
    feuille.append(entete)

    for ligne in lignes:
        feuille.append([str(ligne.get(champ, "") or "") for champ in entete])

    for rang in feuille.iter_rows(min_row=1):
        for cellule in rang:
            cellule.number_format = _FORMAT_TEXTE
            # openpyxl types a string starting with "=" as a formula; a
            # justification or a title must never become one.
            if isinstance(cellule.value, str):
                cellule.data_type = "s"

    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = f"A1:{get_column_letter(len(entete))}{feuille.max_row}"

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(xlsx_path)
    return len(lignes)


def xlsx_vers_csv(xlsx_path: Path, csv_path: Path) -> int:
    """Write a workbook back as the CSV the existing re-import path expects.

    This is the hinge of the return trip (user arbitration of 08/09/2026,
    option A): run_poc006 converts, then runs its current, tested path on the
    result, so profile_store.py - which carries the conflict rules of POC-006
    and POC-009 - is not touched. The intermediate CSV is also an inspectable
    trace of what Excel gave back: on a workbook nobody annotated, it is
    identical to the exported CSV.
    """
    lignes = lire_xlsx(xlsx_path)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=PROFILE_CSV_FIELDS, restval="")
        writer.writeheader()
        writer.writerows(lignes)
    return len(lignes)
